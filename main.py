import os
import asyncio
import logging
import threading
from io import BytesIO

os.environ["CUDA_VISIBLE_DEVICES"] = "-1"
os.environ["MKL_NUM_THREADS"] = "1"
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["TRANSFORMERS_VERBOSITY"] = "error"

import aiohttp
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import CommandStart
from aiogram.client.session.aiohttp import AiohttpSession
from aiogram.fsm.state import StatesGroup, State
from aiogram.fsm.context import FSMContext
from aiogram.utils.keyboard import ReplyKeyboardBuilder
from dotenv import load_dotenv

from pypdf import PdfReader
from docx import Document
import openpyxl

from core.vector_store import VectorStoreManager
from core.llm import ask_llm

from gui import ui_ctrl, create_clean_gui

load_dotenv()
TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN")
if not TELEGRAM_TOKEN:
    raise ValueError("Критическая ошибка: TELEGRAM_TOKEN не найден в .env")

logging.basicConfig(level=logging.ERROR)

dp = Dispatcher()
db = VectorStoreManager()
ui_ctrl.db_reference = db

class NoteStates(StatesGroup):
    waiting_for_note = State()

PROXY_LIST_URL = "https://api.proxyscrape.com/v2/?request=displayproxies&protocol=http&timeout=2000&country=all&ssl=all&anonymity=all"

async def get_working_proxy() -> str | None:
    ui_ctrl.log("SYSTEM", "Scanning proxy nodes...")
    async with aiohttp.ClientSession() as session:
        try:
            async with session.get(PROXY_LIST_URL, timeout=5) as response:
                if response.status != 200: return None
                text = await response.text()
                proxies = [p.strip() for p in text.split("\n") if p.strip()]
        except Exception as e:
            ui_ctrl.log("SYSTEM", f"Proxy provider error: {e}")
            return None

        for proxy_addr in proxies[:15]:
            proxy_url = f"http://{proxy_addr}"
            try:
                async with session.get("https://api.telegram.org", proxy=proxy_url, timeout=2) as test_resp:
                    if test_resp.status in [200, 404]:
                        ui_ctrl.log("SYSTEM", f"Connection verified: {proxy_url}")
                        return proxy_url
            except Exception:
                continue
    return None

async def extract_text_from_document(message: types.Message, bot: Bot) -> tuple[str | None, str]:
    try:
        file_id = message.document.file_id
        file_name = message.document.file_name
        file_ext = file_name.split('.')[-1].lower()
        
        file = await bot.get_file(file_id)
        destination = BytesIO()
        await bot.download(file, destination=destination)
        destination.seek(0)
        
        if file_ext == 'txt':
            return destination.read().decode('utf-8', errors='ignore').strip(), file_name
        elif file_ext == 'pdf':
            reader = PdfReader(destination)
            text = "".join([page.extract_text() or "" for page in reader.pages])
            return text.strip(), file_name
        elif file_ext == 'docx':
            doc = Document(destination)
            text = "\n".join([p.text for p in doc.paragraphs])
            return text.strip(), file_name
        elif file_ext == 'csv':
            return destination.read().decode('utf-8', errors='ignore').strip(), file_name
        elif file_ext == 'xlsx':
            wb = openpyxl.load_workbook(destination, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_str = " | ".join([str(cell) for cell in row if cell is not None])
                    if row_str.strip(): lines.append(row_str)
            return "\n".join(lines).strip(), file_name
            
        return None, file_name
    except Exception as e:
        ui_ctrl.log("ERROR", f"Extraction error: {e}")
        return None, "unknown"

def get_main_keyboard():
    builder = ReplyKeyboardBuilder()
    builder.button(text="📝 Добавить заметку")
    return builder.as_markup(resize_keyboard=True)

@dp.message(CommandStart())
async def cmd_start(message: types.Message, state: FSMContext):
    await state.clear()
    user_id = str(message.from_user.id)
    ui_ctrl.register_user_activity(user_id)
    ui_ctrl.log("BOT", f"Session Init: ID {user_id}")
    
    await message.answer(
        "🧠 <b>Привет! Я твой умный RAG-ассистент для заметок.</b>\n\n"
        "• Чтобы сохранить новую заметку или документ, нажмите кнопку <b>📝 Добавить заметку</b>.\n"
        "• Вы можете <b>переслать</b> мне сообщение или документ из любого чата.\n"
        "• Чтобы найти информацию, просто <b>отправьте мне вопрос</b>.",
        reply_markup=get_main_keyboard(),
        parse_mode="HTML"
    )

@dp.message(F.forward_date | F.forward_from | F.forward_from_chat | F.forward_origin)
async def handle_forwarded_message(message: types.Message, state: FSMContext):
    user_id = str(message.from_user.id)
    ui_ctrl.register_user_activity(user_id)
    text = ""
    
    if message.text:
        text = message.text.strip()
    elif message.caption:
        text = message.caption.strip()
    elif message.document:
        await message.answer("📥 Извлекаю данные...")
        text, f_name = await extract_text_from_document(message, message.bot)
        ui_ctrl.log("DB", f"Parsing file '{f_name}' for {user_id}")

    if not text:
        await message.answer("❌ Формат не поддерживается. Только: TXT, PDF, DOCX, XLSX, CSV.", parse_mode="HTML")
        return

    note_id = f"fwd_{message.message_id}"
    db.save_note(user_id=user_id, note_id=note_id, text=text)
    await state.clear()
    
    try: total_notes = len(db.collection.get()['ids'])
    except: total_notes = 0
    ui_ctrl.refresh_global_stats(os.environ.get("CURRENT_PROXY", "DIRECT"), total_notes, len(ui_ctrl.known_users))
    
    ui_ctrl.log("DB", f"Forwarded chunk committed for {user_id}")
    
    try:
        await message.answer("✅ Данные сохранены в базу знаний!", reply_markup=get_main_keyboard())
    except Exception as e:
        ui_ctrl.log("ERROR", f"Network timeout replying to {user_id}: {e}")

@dp.message(F.text == "📝 Добавить заметку")
async def process_add_note_button(message: types.Message, state: FSMContext):
    await state.set_state(NoteStates.waiting_for_note)
    await message.answer("⏳ <b>Пришлите текст или документ.</b>", parse_mode="HTML")

@dp.message(NoteStates.waiting_for_note)
async def handle_note_input(message: types.Message, state: FSMContext):
    user_id = str(message.from_user.id)
    ui_ctrl.register_user_activity(user_id)
    text = ""

    if message.text:
        text = message.text.strip()
        if text == "📝 Добавить заметку":
            await message.answer("Отправьте текст или файл.")
            return
    elif message.document:
        await message.answer("📥 Обрабатываю документ...")
        text, f_name = await extract_text_from_document(message, message.bot)
        ui_ctrl.log("DB", f"Processing '{f_name}' for {user_id}")
    
    if not text:
        await message.answer("❌ Ошибка чтения.", parse_mode="HTML")
        return

    note_id = str(message.message_id)
    db.save_note(user_id=user_id, note_id=note_id, text=text)
    await state.clear()
    
    try: total_notes = len(db.collection.get()['ids'])
    except: total_notes = 0
    ui_ctrl.refresh_global_stats(os.environ.get("CURRENT_PROXY", "DIRECT"), total_notes, len(ui_ctrl.known_users))
    
    ui_ctrl.log("DB", f"Data indexed for {user_id}")
    
    try:
        await message.answer("✅ Данные проиндексированы!", reply_markup=get_main_keyboard())
    except Exception as e:
        ui_ctrl.log("ERROR", f"Network timeout replying to {user_id}: {e}")

@dp.message()
async def handle_question(message: types.Message):
    user_id = str(message.from_user.id)
    ui_ctrl.register_user_activity(user_id)
    
    query = message.text.strip()
    if not query: return

    ui_ctrl.log("BOT", f"Query from {user_id}: '{query[:30]}...'")
    await message.answer("🔍 Анализирую базу знаний...")
    
    context_list = db.search_notes(user_id=user_id, query=query, limit=6)
    
    if not context_list:
        ui_ctrl.log("BOT", f"No context matches for {user_id}")
        await message.answer("В базе нет схожей информации.")
        return
        
    loop = asyncio.get_running_loop()
    answer = await loop.run_in_executor(None, ask_llm, query, context_list)
    
    ui_ctrl.log("LLM", f"Response ready for {user_id}")
    await message.answer(answer)

async def run_bot():
    working_proxy = await get_working_proxy()
    
    try:
        all_data = db.collection.get()
        total_d = len(all_data.get('ids', []))
        if all_data.get('metadatas'):
            for meta in all_data['metadatas']:
                if meta and 'user_id' in meta:
                    ui_ctrl.register_user_activity(str(meta['user_id']))
    except:
        total_d = 0
    
    ui_ctrl.refresh_global_stats(working_proxy if working_proxy else "DIRECT", total_d, len(ui_ctrl.known_users))
    os.environ["CURRENT_PROXY"] = working_proxy if working_proxy else "DIRECT"

    bot = None
    try:
        if working_proxy:
            ui_ctrl.log("SYSTEM", f"Starting bot via proxy: {working_proxy}")
            session = AiohttpSession(proxy=working_proxy)
            bot = Bot(token=TELEGRAM_TOKEN, session=session)
            await bot.get_me()
        else:
            raise Exception("No proxy available")
            
    except Exception as e:
        ui_ctrl.log("SYSTEM", f"Proxy connection failed: {e}. Switching to DIRECT mode.")
        bot = Bot(token=TELEGRAM_TOKEN)
        os.environ["CURRENT_PROXY"] = "DIRECT"

    ui_ctrl.log("SYSTEM", "Core Operational.")
    try:
        await dp.start_polling(bot)
    finally:
        await bot.session.close()

def start_bot_thread():
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    loop.run_until_complete(run_bot())

if __name__ == "__main__":
    create_clean_gui(start_bot_thread)